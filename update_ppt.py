"""
Python script to update 'ppt sample.pptx' with EVAL AI project content using clean bullet points.
"""

import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def update_presentation(pptx_path, output_path):
    prs = Presentation(pptx_path)
    
    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides[0]
    for shape in slide1.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text.strip()
        if "STOCK MARKET" in text or "DATA ANALYSIS" in text or "AUTONOMOUS RESUME" in text:
            shape.text_frame.clear()
            p = shape.text_frame.add_paragraph()
            p.text = "EVAL AI : AUTONOMOUS RESUME ANALYZER &\nAI INTERVIEW AGENT"
            p.font.bold = True
            p.font.size = Pt(22)
                
    # -------------------------------------------------------------
    # SLIDE 2: Agenda
    # -------------------------------------------------------------
    slide2 = prs.slides[1]
    for shape in slide2.shapes:
        if shape.has_text_frame and ("Agenda" in shape.name or "Content Placeholder" in shape.name or "1." in shape.text or "Title" in shape.text):
            shape.text_frame.clear()
            agenda_items = [
                "1. Project Title & Vision",
                "2. Abstract",
                "3. Introduction & Problem Statement",
                "4. System Architecture & Workflow",
                "5. Literature Review & Methodology",
                "6. Technology Stack & Tools Used",
                "7. Design & Implementation (ML + GenAI)",
                "8. Key Features & Industry Use Cases",
                "9. Model Evaluation & Performance Metrics",
                "10. Conclusion & Future Scope"
            ]
            for item in agenda_items:
                p = shape.text_frame.add_paragraph()
                p.text = item
                p.font.size = Pt(13)
                p.space_after = Pt(5)

    # -------------------------------------------------------------
    # SLIDE 3: Title Slide 2
    # -------------------------------------------------------------
    slide3 = prs.slides[2]
    for shape in slide3.shapes:
        if shape.has_text_frame and ("STOCK" in shape.text or "PREDICTION" in shape.text or "RESUME" in shape.text or shape.name == "TextBox 2"):
            shape.text_frame.clear()
            p = shape.text_frame.add_paragraph()
            p.text = "EVAL AI: RESUME ANALYZER, SUPERVISED ML\nSUITABILITY PREDICTOR & AI INTERVIEW AGENT"
            p.font.bold = True
            p.font.size = Pt(22)

    # -------------------------------------------------------------
    # SLIDE 4: Abstract
    # -------------------------------------------------------------
    slide4 = prs.slides[3]
    for shape in slide4.shapes:
        if shape.has_text_frame and ("Content Placeholder" in shape.name or "EVAL AI is" in shape.text or "The system is" in shape.text):
            shape.text_frame.clear()
            p1 = shape.text_frame.add_paragraph()
            p1.text = (
                "EVAL AI is an enterprise-grade autonomous recruitment and talent evaluation system "
                "integrating Supervised Machine Learning and Generative AI. The platform performs multi-format "
                "resume parsing (PDF, DOCX, ZIP batch), extracts structured skill and experience vectors, and "
                "utilizes a trained Random Forest Classifier (120 estimators, max depth=9) to compute deterministic "
                "candidate-to-job suitability classifications and probability scores."
            )
            p1.font.size = Pt(13)
            p1.space_after = Pt(8)
            
            p2 = shape.text_frame.add_paragraph()
            p2.text = (
                "Suitable candidates are routed to an adaptive 4-round interview simulation (Quantitative Aptitude, "
                "Domain Technical MCQs, Architecture Scenarios, and Spoken Voice Assessment) powered by Groq LLMs "
                "with real-time MediaPipe biometric proctoring, automated scoring, and comprehensive PDF dossier generation."
            )
            p2.font.size = Pt(13)

    # -------------------------------------------------------------
    # SLIDE 5: Introduction
    # -------------------------------------------------------------
    slide5 = prs.slides[4]
    for shape in slide5.shapes:
        if shape.has_text_frame and ("Content Placeholder" in shape.name or "In contemporary" in shape.text or "data-driven" in shape.text):
            shape.text_frame.clear()
            p1 = shape.text_frame.add_paragraph()
            p1.text = (
                "In contemporary talent acquisition, screening high-volume applicant pipelines manually or with "
                "static keyword-based ATS filters results in significant false positives, high operational latency, "
                "and lack of objective skill verification. Furthermore, relying purely on GenAI for scoring introduces "
                "hallucinations and non-deterministic decision boundaries."
            )
            p1.font.size = Pt(12)
            p1.space_after = Pt(6)
            
            p2 = shape.text_frame.add_paragraph()
            p2.text = "EVAL AI solves this through a dual-engine architecture:"
            p2.font.bold = True
            p2.font.size = Pt(12)
            p2.space_after = Pt(4)
            
            p3 = shape.text_frame.add_paragraph()
            p3.text = "- Supervised ML Classifier: Evaluates an 11-dimensional numerical feature vector (skill match ratio, project density, experience, certifications) for transparent suitability prediction."
            p3.font.size = Pt(11.5)
            p3.space_after = Pt(3)

            p4 = shape.text_frame.add_paragraph()
            p4.text = "- GenAI Interview Agent: Conducts dynamic multi-round technical and behavioral assessments calibrated to candidate competency profiles and identified skill gaps."
            p4.font.size = Pt(11.5)

    # -------------------------------------------------------------
    # SLIDE 6: System Overview
    # -------------------------------------------------------------
    slide6 = prs.slides[5]
    for shape in slide6.shapes:
        if shape.has_text_frame and ("Rectangle 1" in shape.name or "proposed system" in shape.text):
            shape.text_frame.clear()
            p0 = shape.text_frame.add_paragraph()
            p0.text = "The proposed system follows a modular, reactive multi-tier architecture:"
            p0.font.bold = True
            p0.font.size = Pt(12)
            p0.space_after = Pt(4)
            
            points = [
                "1. Document Ingestion Tier: High-throughput batch ZIP/PDF extraction with OCR fallback and structured JSON tokenization.",
                "2. Supervised ML Tier (FastAPI): 11-feature mathematical extraction pipeline invoking a trained Random Forest model (89.75% test accuracy) for deterministic suitability probability.",
                "3. GenAI Assessment Engine: 4-round dynamic interview engine (Quantitative Aptitude, Domain MCQs, Architecture Scenarios, and Web Speech API Voice Evaluation).",
                "4. Anti-Malpractice Proctoring: MediaPipe facial biometric tracking, yaw/pitch pose estimation, multi-face alerts, and tab-switch detection.",
                "5. Enterprise Core (Spring Boot + React): Java 21 JPA persistence, candidate ranking, and instant PDF dossier synthesis."
            ]
            for pt in points:
                p = shape.text_frame.add_paragraph()
                p.text = pt
                p.font.size = Pt(11)
                p.space_after = Pt(3)

    # -------------------------------------------------------------
    # SLIDE 7: Literature Review
    # -------------------------------------------------------------
    slide7 = prs.slides[6]
    for shape in slide7.shapes:
        if shape.has_text_frame and ("Content Placeholder" in shape.name or "Resume Feature Extraction" in shape.text or "foundational" in shape.text):
            shape.text_frame.clear()
            reviews = [
                ("1. Resume Feature Extraction & NLP:", " Traditional TF-IDF keyword matching fails to capture semantic synonyms and contextual depth. Modern entity extraction combined with structured taxonomies drastically improves feature vector representations."),
                ("2. Ensemble Machine Learning in Hiring:", " Supervised ensemble algorithms (such as Random Forest Classifiers) provide non-linear decision boundaries, robust handling of multi-modal features, and Mean Decrease in Impurity (MDI) feature importance rankings without black-box opacity."),
                ("3. Adaptive Conversational AI in Interviewing:", " Large Language Models (LLMs) enable contextual, multi-turn technical dialogues and automated code/architecture critique, surpassing static questionnaire banks.")
            ]
            for heading, body in reviews:
                p = shape.text_frame.add_paragraph()
                p.text = heading + body
                p.font.size = Pt(11.5)
                p.space_after = Pt(5)

    # -------------------------------------------------------------
    # SLIDE 8: Tools Used
    # -------------------------------------------------------------
    slide8 = prs.slides[7]
    for shape in slide8.shapes:
        if shape.has_text_frame and ("TextBox 6" in shape.name or "FRONTEND:" in shape.text):
            shape.text_frame.clear()
            tools_list = [
                "• FRONTEND: React.js (Vite), Royal Dark Glassmorphism UI, Lucide Icons, Canvas Radar Charts.",
                "• BACKEND: Spring Boot 3 (Java 21), Spring Data JPA, Oracle / H2 Database, Apache PDFBox, Apache POI.",
                "• MACHINE LEARNING: Python 3.11, scikit-learn (RandomForestClassifier), Pandas, NumPy, Joblib, FastAPI, Uvicorn.",
                "• GENAI & NLP: Groq Cloud API (Llama 3.3, Qwen 2.5), Web Speech API (Voice-to-Text).",
                "• AI PROCTORING: MediaPipe Face Mesh, Live WebCam Biometrics & Tab Tracking."
            ]
            for t in tools_list:
                p = shape.text_frame.add_paragraph()
                p.text = t
                p.font.size = Pt(12)
                p.space_after = Pt(5)

    # -------------------------------------------------------------
    # SLIDE 9: Design and Implementation
    # -------------------------------------------------------------
    slide9 = prs.slides[8]
    for shape in list(slide9.shapes):
        if shape.has_text_frame and shape.name == "Title 1":
            shape.text = "DESIGN AND IMPLEMENTATION ARCHITECTURE"
        elif shape.has_text_frame and shape.name != "Title 1" and "Phase" in shape.text:
            shape.text_frame.clear()
            
    # Find existing or add textbox
    tb = None
    for shape in slide9.shapes:
        if shape.has_text_frame and shape.name not in ["Title 1", "Date Placeholder 3", "Footer Placeholder 4", "Slide Number Placeholder 5"]:
            tb = shape
            break
            
    if not tb:
        tb = slide9.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(4.8))
        
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    
    stages = [
        "Phase 1: Resume Parsing & Feature Extraction",
        "  - Multi-format PDF/DOCX parsed into structured JSON profile (skills, projects, experience).",
        "  - Extracted 11-dimensional numerical feature vector.",
        "",
        "Phase 2: Supervised ML Suitability Classification",
        "  - RandomForestClassifier (n_estimators=120, max_depth=9, random_state=42).",
        "  - Returns binary classification (Suitable / Not Suitable) and probability score.",
        "  - Computes MDI feature importance rankings and identifies missing skill gaps.",
        "",
        "Phase 3: 4-Round Adaptive AI Interview Simulation",
        "  - Round 1: Quantitative Aptitude & Logical Reasoning.",
        "  - Round 2: Domain Technical MCQs (Contextualized by ML suitability output).",
        "  - Round 3: Practical Architecture & Scenario Design.",
        "  - Round 4: Voice Speech & Linguistic Competency Evaluation.",
        "",
        "Phase 4: Multi-Tier Assessment Dossier & Ranking",
        "  - Real-time biometric proctoring tracking (MediaPipe Face Mesh).",
        "  - Instant candidate dossier generation & PDF report synthesis."
    ]
    for s in stages:
        p = tf.add_paragraph()
        p.text = s
        if s.startswith("Phase"):
            p.font.bold = True
            p.font.size = Pt(12)
            p.font.color.rgb = RGBColor(212, 175, 55) # Royal Gold
        else:
            p.font.size = Pt(10.5)

    # -------------------------------------------------------------
    # SLIDE 10: Use Cases
    # -------------------------------------------------------------
    slide10 = prs.slides[9]
    for shape in slide10.shapes:
        if shape.has_text_frame and ("Rectangle 2" in shape.name or "Campus" in shape.text or "Marketing" in shape.text):
            shape.text_frame.clear()
            use_cases = [
                "• Automated Campus Recruitment & Bulk ZIP Batch Processing",
                "• Objective Role-Specific Candidate Suitability Shortlisting",
                "• Supervised ML Skill Gap Analysis & Technical Profiling",
                "• AI-Proctored Remote Screening & Integrity Verification",
                "• Adaptive Multi-Round Technical Interview Simulation",
                "• Voice Articulation & Linguistic Proficiency Assessment",
                "• Instant Candidate PDF Dossier & Executive Report Generation",
                "• Real-Time Recruiter Analytics & Cohort Leaderboard Tracking"
            ]
            for uc in use_cases:
                p = shape.text_frame.add_paragraph()
                p.text = uc
                p.font.size = Pt(12.5)
                p.space_after = Pt(5)

    # -------------------------------------------------------------
    # SLIDE 11: Conclusion
    # -------------------------------------------------------------
    slide11 = prs.slides[10]
    for shape in slide11.shapes:
        if shape.has_text_frame and ("Rectangle 2" in shape.name or "EVAL AI successfully" in shape.text or "The AI-Based" in shape.text):
            shape.text_frame.clear()
            p0 = shape.text_frame.add_paragraph()
            p0.text = "EVAL AI successfully establishes a scalable, fair, and mathematically grounded recruitment automation platform."
            p0.font.bold = True
            p0.font.size = Pt(12.5)
            p0.space_after = Pt(6)
            
            conclusions = [
                "• Separation of Concerns: Machine Learning handles deterministic numerical prediction and feature importances, while Generative AI powers interactive interview simulation and language critique.",
                "• High Evaluation Precision: The Random Forest model delivers 89.75% test accuracy, 92.61% precision, and 90.17% F1-score on realistic hiring benchmarks.",
                "• Integrity-First Evaluation: MediaPipe live face tracking and tab monitoring ensure fair, cheat-resilient remote assessments.",
                "• End-to-End Automation: Reduces technical hiring triage cycles from days to minutes while delivering rich, verifiable candidate dossiers."
            ]
            for c in conclusions:
                p = shape.text_frame.add_paragraph()
                p.text = c
                p.font.size = Pt(11.5)
                p.space_after = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 12: Thank You
    # -------------------------------------------------------------
    slide12 = prs.slides[11]
    for shape in slide12.shapes:
        if shape.has_text_frame and ("THANK YOU" in shape.text or shape.name == "TextBox 2"):
            shape.text_frame.clear()
            p1 = shape.text_frame.add_paragraph()
            p1.text = "THANK YOU"
            p1.font.bold = True
            p1.font.size = Pt(36)
            p1.space_after = Pt(10)
            
            p2 = shape.text_frame.add_paragraph()
            p2.text = "EVAL AI: Autonomous Talent Assessment & ML Prediction Platform"
            p2.font.size = Pt(16)
            p2.space_after = Pt(6)

            p3 = shape.text_frame.add_paragraph()
            p3.text = "Questions & Discussion"
            p3.font.size = Pt(14)
            p3.font.italic = True

    prs.save(output_path)
    print(f"[OK] Successfully saved: {output_path}")

if __name__ == "__main__":
    src_file = r"D:\KIT Hackathon\AI-ResumeAnalyzer\ppt sample.pptx"
    out_file = r"D:\KIT Hackathon\AI-ResumeAnalyzer\EVAL_AI_ResumeAnalyzer_Presentation.pptx"
    update_presentation(src_file, out_file)
    update_presentation(src_file, src_file)
